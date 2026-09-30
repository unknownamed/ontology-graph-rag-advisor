"""Local document extraction. Every parsed course is an untrusted candidate."""
from __future__ import annotations

import hashlib
import io
import re
import sys
import tempfile
import zipfile
from collections import Counter
from contextlib import closing
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
MAX_BYTES = 10_000_000
COURSE_CODE = re.compile(r"(?<![A-Z0-9])([A-Z]{3}\d{4})(?![A-Z0-9])", re.I)
SEMESTER = re.compile(r"(?<!\d)(20\d{2})\s*(?:년\s*)?[-./]?\s*([12])\s*학기|(?<!\d)(20\d{2})-([12])(?!\d)")
CREDIT = re.compile(r"(?<!\d)(\d{1,2})\s*(?:학점|credits?)(?!\w)", re.I)
CLASSIFICATION = re.compile(r"전공필수|전공선택|교양필수|교양선택|기초교양|균형교양|확대교양|일반선택|자유선택")
COMPLETION = re.compile(r"미이수|미취득|수강중|낙제|이수|취득|완료")
GRADE = re.compile(r"(?:[A-D][+0-]?|F|S|U|P)", re.I)
_OCR_READER = None


def _pdf_lines(data: bytes) -> tuple[list[tuple[str, str]], list[str], dict[str, float], int, list[int]]:
    if not data.startswith(b"%PDF"):
        raise ValueError("PDF signature is invalid")
    import pdfplumber
    lines: list[tuple[str, str]] = []
    warnings: list[str] = []
    scores: dict[str, float] = {}
    unreadable_pages: list[int] = []
    with pdfplumber.open(io.BytesIO(data)) as document:
        if len(document.pages) > 40:
            raise ValueError("PDF exceeds the 40-page upload limit")
        for number, page in enumerate(document.pages, 1):
            content = page.extract_text() or ""
            if not content.strip():
                if number <= 10:
                    try:
                        scanned, _, confidence = _ocr_image(page.to_image(resolution=150).original,
                                                             f"PDF page {number} OCR")
                        lines.extend(scanned)
                        scores.update(confidence)
                        warnings.append(f"PDF {number}쪽은 로컬 OCR 후보로 처리했습니다. 사용자 확인이 필요합니다.")
                        if not scanned:
                            unreadable_pages.append(number)
                    except ValueError as exc:
                        warnings.append(f"PDF {number}쪽 OCR 불가: {exc}")
                        unreadable_pages.append(number)
                else:
                    warnings.append(f"PDF {number}쪽은 텍스트가 없고 OCR 10쪽 한도를 초과했습니다.")
                    unreadable_pages.append(number)
            table_rows: list[tuple[str, str, str]] = []
            for table_no, table in enumerate(page.extract_tables(), 1):
                for row_no, row in enumerate(table, 1):
                    cells = [" ".join((cell or "").split()) for cell in row]
                    code_cells = [cell.upper() for cell in cells if COURSE_CODE.fullmatch(cell)]
                    if (len(code_cells) <= 1 and len(cells) >= 6 and
                            re.fullmatch(r"20\d{2}", cells[0]) and
                            cells[1] in {"1", "2", "하계", "동계"}):
                        table_rows.append((f"PDF page {number}, table {table_no}, row {row_no}",
                                           " | ".join(cells), code_cells[0] if code_cells else None))
            table_code_counts = Counter(code for _, _, code in table_rows if code)
            for index, value in enumerate(content.splitlines(), 1):
                code_match = COURSE_CODE.search(value)
                if (code_match and table_code_counts[code_match.group(1).upper()] and
                        re.match(r"^\s*20\d{2}\s+(?:[12]|하계|동계)\s+", value)):
                    table_code_counts[code_match.group(1).upper()] -= 1
                    continue
                lines.append((f"PDF page {number}, line {index}", value))
            lines.extend((location, value) for location, value, _ in table_rows)
            if sum(len(value) for _, value in lines) > 2_000_000:
                raise ValueError("Extracted PDF text exceeds the 2,000,000-character limit")
    return lines, warnings, scores, len(document.pages), unreadable_pages


def _docx_lines(data: bytes) -> tuple[list[tuple[str, str]], list[str]]:
    if not data.startswith(b"PK"):
        raise ValueError("DOCX signature is invalid")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        files = archive.infolist()
        if len(files) > 1000 or sum(item.file_size for item in files) > 30_000_000:
            raise ValueError("DOCX decompressed content exceeds the limit")
    from docx import Document
    document = Document(io.BytesIO(data))
    lines = [(f"paragraph {i}", p.text) for i, p in enumerate(document.paragraphs, 1)]
    for table_index, table in enumerate(document.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            lines.append((f"table {table_index}, row {row_index}", " | ".join(cell.text for cell in row.cells)))
    if sum(len(value) for _, value in lines) > 2_000_000:
        raise ValueError("Extracted DOCX text exceeds the limit")
    return lines, []


def _hwp_lines(data: bytes) -> tuple[list[tuple[str, str]], list[str]]:
    if not data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        raise ValueError("Only HWP 5 OLE documents are supported")
    dependency_dir = ROOT / ".deps"
    if str(dependency_dir) not in sys.path:
        sys.path.insert(0, str(dependency_dir))
    try:
        from hwp5.hwp5txt import TextTransform
        from hwp5.xmlmodel import Hwp5File
    except ImportError as exc:
        raise ValueError("HWP parser unavailable; run the documented pyhwp2 setup") from exc
    with tempfile.TemporaryDirectory(prefix="curriculum-hwp-") as folder:
        path = Path(folder) / "upload.hwp"
        path.write_bytes(data)
        output = io.BytesIO()
        with closing(Hwp5File(str(path))) as document:
            TextTransform().transform_hwp5_to_text(document, output)
        content = output.getvalue().decode("utf-8")
    if len(content) > 2_000_000:
        raise ValueError("Extracted HWP text exceeds the limit")
    return [(f"text line {i}", line) for i, line in enumerate(content.splitlines(), 1)], []


def _image_lines(data: bytes) -> tuple[list[tuple[str, str]], list[str], dict[str, float]]:
    from PIL import Image
    from PIL import UnidentifiedImageError
    try:
        image = Image.open(io.BytesIO(data))
        if image.width * image.height > 20_000_000:
            raise ValueError("Image exceeds the 20-megapixel OCR limit")
        image.load()
        image = image.convert("RGB")
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError("Invalid image upload") from exc
    return _ocr_image(image, "OCR")


def _ocr_image(image, prefix: str) -> tuple[list[tuple[str, str]], list[str], dict[str, float]]:
    dependency_dir = ROOT / ".ocrdeps"
    if str(dependency_dir) not in sys.path:
        sys.path.insert(0, str(dependency_dir))
    try:
        import easyocr
        import numpy as np
    except ImportError as exc:
        raise ValueError("Korean OCR unavailable; run the documented EasyOCR setup") from exc
    global _OCR_READER
    if _OCR_READER is None:
        models = ROOT / ".ocrmodels"
        if not (models / "craft_mlt_25k.pth").exists() or not (models / "korean_g2.pth").exists():
            raise ValueError("Local Korean OCR model files are missing")
        _OCR_READER = easyocr.Reader(["ko", "en"], gpu=False,
                                     model_storage_directory=str(models), download_enabled=False,
                                     verbose=False)
    result = _OCR_READER.readtext(np.asarray(image), detail=1)
    lines = [(f"{prefix} line {i}", text) for i, (_, text, _) in enumerate(result, 1)]
    scores = {f"{prefix} line {i}": float(confidence) for i, (_, _, confidence) in enumerate(result, 1)}
    return lines, [], scores


def _compact(value: str | None) -> str:
    return "".join((value or "").split()).casefold()


def _course_name_key(value: str | None) -> str:
    """Ignore only a trailing delivery-mode note, preserving the uploaded spelling."""
    return re.sub(r"\(인강\)$", "", _compact(value))


def _raw_course_values(text: str, code_match: re.Match | None) -> dict:
    cells = [part.strip() for part in text.split("|")] if "|" in text else []
    raw_name = None
    code_cell = None
    if code_match:
        if cells:
            code_cell = next((i for i, cell in enumerate(cells) if COURSE_CODE.search(cell)), None)
            if code_cell is not None and code_cell + 1 < len(cells):
                following = cells[code_cell + 1]
                if following and not CLASSIFICATION.fullmatch(following):
                    raw_name = following
        else:
            following = text[code_match.end():]
            following = re.split(r"\d{1,2}\s*학점|20\d{2}\s*[-./년]?\s*[12]\s*학기|전공필수|전공선택|기초교양|균형교양|확대교양|교양필수|교양선택|일반선택|자유선택|\b(?:A\+|B\+|C\+|A0|B0|C0|P|F)\b|이수|취득|완료", following, maxsplit=1)[0].strip(" -:,;")
            if re.fullmatch(r"[가-힣A-Za-z][가-힣A-Za-z\s·()]{1,45}", following):
                raw_name = following
    else:
        raw_name = next((cell for cell in cells if re.fullmatch(r"[가-힣A-Za-z][가-힣A-Za-z\s·()]{1,45}", cell)
                         and not CLASSIFICATION.fullmatch(cell)), None)
    credit = CREDIT.search(text)
    raw_credit = int(credit.group(1)) if credit else None
    if raw_credit is None and code_cell is not None and code_cell + 2 < len(cells):
        credit_cell = cells[code_cell + 2]
        if re.fullmatch(r"\d{1,2}", credit_cell) and int(credit_cell) <= 30:
            raw_credit = int(credit_cell)
    if (raw_credit is None and len(cells) == 7 and
            re.fullmatch(r"20\d{2}", cells[0]) and cells[1] in {"1", "2", "하계", "동계"} and
            not cells[3] and re.fullmatch(r"\d{1,2}", cells[5]) and int(cells[5]) <= 30):
        # A blank code cell does not shift the credit column in a seven-column transcript table.
        raw_credit = int(cells[5])
    if raw_credit is None and cells:
        numbers = [int(cell) for cell in cells if re.fullmatch(r"\d{1,2}", cell) and int(cell) <= 30]
        raw_credit = numbers[0] if len(numbers) == 1 else None
    term = SEMESTER.search(text)
    semester = f"{term.group(1) or term.group(3)}-{term.group(2) or term.group(4)}" if term else None
    if not semester and cells and re.fullmatch(r"20\d{2}-(?:SUMMER|WINTER)", cells[0], re.I):
        semester = cells[0].upper()
    if not semester and len(cells) >= 2 and re.fullmatch(r"20\d{2}", cells[0]):
        suffix = {"1": "1", "2": "2", "하계": "SUMMER", "동계": "WINTER"}.get(cells[1])
        if suffix:
            semester = f"{cells[0]}-{suffix}"
    if not semester and not cells and code_match:
        prefix = text[:code_match.start()].split(maxsplit=2)
        if len(prefix) >= 2 and re.fullmatch(r"20\d{2}", prefix[0]):
            suffix = {"1": "1", "2": "2", "하계": "SUMMER", "동계": "WINTER"}.get(prefix[1])
            if suffix:
                semester = f"{prefix[0]}-{suffix}"
        trailing = re.fullmatch(r"(.+?)\s+(\d{1,2})\s+([A-D][+0-]?|F|S|U|P)", text[code_match.end():].strip(), re.I)
        if trailing:
            raw_name = trailing.group(1)
            raw_credit = int(trailing.group(2))
    classification = CLASSIFICATION.search(text)
    raw_classification = (cells[code_cell - 1] if code_cell is not None and code_cell > 0
                          and CLASSIFICATION.search(cells[code_cell - 1]) else
                          classification.group(0) if classification else None)
    grade = cells[-1] if cells and GRADE.fullmatch(cells[-1]) else None
    if grade is None and code_match and not cells:
        last = text.rsplit(maxsplit=1)[-1]
        grade = last if GRADE.fullmatch(last) else None
    completion = COMPLETION.search(text)
    result = {"course_name": raw_name, "earned_credits": raw_credit, "semester": semester,
              "classification": raw_classification,
              "completion_marker": completion.group(0) if completion else grade}
    if grade:
        result["grade"] = grade
    return result


def _student_fields(lines: list[tuple[str, str]]) -> dict:
    patterns = {
        "student_id": re.compile(r"(?:학번|학생번호)\s*[:：]?\s*([0-9]{6,12})(?!\d)"),
        "admission_year": re.compile(r"(?:입학연도|입학년도)\s*[:：]?\s*(20\d{2})(?!\d)"),
        "department": re.compile(r"(?:학과|소속)\s*[:：]?\s*([가-힣A-Za-z ]{2,30})(?=\s*[|,;]|$)"),
    }
    result = {}
    for name, pattern in patterns.items():
        matches = [{"raw": match.group(1).strip(), "location": location}
                   for location, text in lines for match in pattern.finditer(text)]
        if name == "department":
            matches = [item for item in matches if not re.search(r"학번|이름|성명|학생번호", item["raw"])]
        if name in {"student_id", "department"}:
            for index, (location, text) in enumerate(lines[:-1]):
                if re.search(r"학과\s+학번\s+(?:이름|성명)", text):
                    next_location, next_text = lines[index + 1]
                    identity = re.search(r"(?:^|\s)([가-힣A-Za-z]+(?:학과|학부))\s+(\d{6,12})(?=\s|$)", next_text)
                    if identity:
                        matches.append({"raw": identity.group(2 if name == "student_id" else 1),
                                        "location": next_location})
        unique = {_compact(item["raw"]) for item in matches}
        result[name] = {"observations": matches, "status": "AMBIGUOUS" if len(unique) > 1 else
                        "UNVERIFIED" if matches else "MISSING"}
    return result


def _candidates(lines: list[tuple[str, str]], catalog: dict, source_id: str,
                ocr_scores: dict[str, float] | None = None) -> list[dict]:
    known = {course["course_id"]: course for course in catalog["courses"]}
    names: dict[str, list[dict]] = {}
    for course in catalog["courses"]:
        if course["verification_status"] == "VERIFIED":
            names.setdefault(_compact(course["name"]), []).append(course)
    found: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for location, raw in lines:
        text = " ".join(raw.split())
        matches = list(COURSE_CODE.finditer(text))
        if not matches and "|" in text:
            cells = [part.strip() for part in text.split("|")]
            name_cell = next((cell for cell in cells if _compact(cell) in names), None)
            if name_cell:
                matches = [None]
        for match in matches:
            raw_code = match.group(1) if match else None
            code = raw_code.upper() if raw_code else None
            key = (location, code or name_cell)
            if key in seen:
                continue
            seen.add(key)
            values = _raw_course_values(text, match)
            if match is None:
                values["course_name"] = name_cell
            course = known.get(code) if code else None
            name_matches = names.get(_compact(values["course_name"]), []) if values["course_name"] else []
            status = ("UNRESOLVED" if code and not course else "AMBIGUOUS" if not code or
                      (values["course_name"] and _course_name_key(values["course_name"]) !=
                       _course_name_key(course["name"]))
                      else "RESOLVED_TO_CATALOG")
            found.append({"candidate_id": f"UPLOAD-{source_id[:12]}-{len(found)+1}",
                          "course_id": code, "raw_values": {"course_code": raw_code, **values},
                          "catalog_course_id": course["course_id"] if course else None,
                          "catalog_name": course["name"] if course else None,
                          "catalog_credits": course["catalog_credits"] if course else None,
                          "catalog_classification": course["classification"] if course else None,
                          "catalog_source": course["source"] if course else None,
                          "catalog_verification_status": course["verification_status"] if course else "MISSING",
                          "possible_course_ids": sorted(item["course_id"] for item in name_matches),
                          "resolution_status": status, "usable_for_decision": False,
                          "location": location, "snippet": text[:240],
                          "extraction_status": "UNVERIFIED", "completion_status": "UNKNOWN",
                          "confidence": "OCR_CANDIDATE" if ocr_scores and location in ocr_scores else "TEXT_CANDIDATE",
                          "ocr_confidence_score": ocr_scores.get(location) if ocr_scores else None})
    repeated = {}
    for candidate in found:
        if candidate["course_id"]:
            repeated.setdefault(candidate["course_id"], []).append(candidate)
    for group in repeated.values():
        if len(group) > 1:
            for candidate in group:
                candidate["resolution_status"] = "DUPLICATE_CANDIDATE"
                candidate["duplicate_candidate_ids"] = [other["candidate_id"] for other in group if other is not candidate]
    return found


def extract_upload(filename: str, data: bytes, catalog: dict) -> dict:
    """Extract candidate mentions; never declare that a student completed a course."""
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Upload must contain 1 to 10,000,000 bytes")
    kind = Path(filename).suffix.lower()
    ocr_scores: dict[str, float] = {}
    page_count = None
    unreadable_pages: list[int] = []
    try:
        if kind == ".pdf":
            lines, warnings, ocr_scores, page_count, unreadable_pages = _pdf_lines(data)
        elif kind == ".docx":
            lines, warnings = _docx_lines(data)
        elif kind == ".hwp":
            lines, warnings = _hwp_lines(data)
        elif kind in {".png", ".jpg", ".jpeg", ".webp"}:
            lines, warnings, ocr_scores = _image_lines(data)
        else:
            raise ValueError("Supported extensions: PDF, DOCX, HWP 5, PNG, JPG, WEBP")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not parse {kind or 'file'} upload") from exc
    source_id = hashlib.sha256(data).hexdigest()
    candidates = _candidates(lines, catalog, source_id, ocr_scores)
    if not lines:
        warnings.append("추출 가능한 텍스트가 없습니다.")
    return {"source_id": source_id, "filename": Path(filename).name,
            "file_type": kind[1:].upper(), "line_count": len(lines),
            "candidate_count": len(candidates), "candidates": candidates, "warnings": warnings,
            "student_fields": _student_fields(lines),
            "document_coverage": {"page_count": page_count, "unreadable_pages": unreadable_pages,
                                  "detected_semesters": sorted({c["raw_values"]["semester"] for c in candidates
                                                                if c["raw_values"]["semester"]}),
                                  "missing_pages": None, "missing_semesters": None,
                                  "completion_coverage": "UNKNOWN"},
            "requires_user_confirmation": True,
            "catalog_courses": [{"course_id": course["course_id"], "name": course["name"],
                                 "credits": course["catalog_credits"]} for course in catalog["courses"]
                                if course["verification_status"] == "VERIFIED"]}


def normalize_upload(extracted: dict, review: dict, catalog: dict) -> dict:
    """Build a partial StudentState from reviewed observations, never a complete transcript claim."""
    if not isinstance(review, dict) or not isinstance(review.get("records"), list):
        raise ValueError("Upload review records are required")
    candidates = {item["candidate_id"]: item for item in extracted["candidates"]}
    submitted = {}
    for row in review["records"]:
        if not isinstance(row, dict) or row.get("candidate_id") not in candidates or row["candidate_id"] in submitted:
            raise ValueError("Unknown or repeated upload candidate ID")
        submitted[row["candidate_id"]] = row
    known = {course["course_id"]: course for course in catalog["courses"]
             if course["verification_status"] == "VERIFIED"}
    prepared = []
    for candidate in extracted["candidates"]:
        row = submitted.get(candidate["candidate_id"], {})
        raw = candidate["raw_values"]
        code = str(row.get("course_code") or raw["course_code"] or "").strip().upper()
        if not COURSE_CODE.fullmatch(code):
            code = ""
        course = known.get(code)
        name = str(row.get("course_name") or raw["course_name"] or "").strip()
        raw_credits = row.get("earned_credits", raw["earned_credits"])
        try:
            credits = int(raw_credits) if raw_credits is not None and str(raw_credits).strip() != "" else None
        except (TypeError, ValueError) as exc:
            raise ValueError("Reviewed earned credits must be an integer or blank") from exc
        if credits is not None and not 0 <= credits <= 30:
            raise ValueError("Reviewed earned credits are outside 0–30")
        completion = row.get("completion_status", "UNKNOWN")
        if completion not in {"COMPLETED", "IN_PROGRESS", "FAILED", "UNKNOWN"}:
            raise ValueError("Reviewed completion status is invalid")
        semester = str(row.get("semester") or raw["semester"] or "").strip() or None
        if semester is not None and not re.fullmatch(r"20\d{2}-(?:[12]|SUMMER|WINTER)", semester):
            raise ValueError("Reviewed semester must be YYYY-1, YYYY-2, YYYY-SUMMER, or YYYY-WINTER")
        classification = str(row.get("classification") or raw["classification"] or "").strip() or None
        prepared.append({"candidate": candidate, "review": row, "code": code, "course": course,
                         "name": name, "credits": credits, "completion": completion,
                         "semester": semester, "classification": classification})
    code_counts = {}
    for item in prepared:
        if item["code"]:
            code_counts[item["code"]] = code_counts.get(item["code"], 0) + 1
    classification_names = {"MAJOR_REQUIRED": "전공필수", "MAJOR_ELECTIVE": "전공선택",
                            "GENERAL_BASIC": "기초교양", "GENERAL_BALANCED": "균형교양",
                            "GENERAL_EXPANDED": "확대교양"}
    attempts, records, attention = [], [], []
    for item in prepared:
        candidate, row, course = item["candidate"], item["review"], item["course"]
        reasons = []
        if not course:
            reasons.append("OFFICIAL_COURSE_NOT_RESOLVED")
        if not candidate["raw_values"]["course_code"]:
            reasons.append("COURSE_CODE_NOT_IN_UPLOAD")
        if not item["name"]:
            reasons.append("COURSE_NAME_MISSING")
        elif course and _course_name_key(item["name"]) != _course_name_key(course["name"]):
            reasons.append("COURSE_NAME_CONFLICT")
        if not candidate["raw_values"]["course_name"]:
            reasons.append("COURSE_NAME_NOT_IN_UPLOAD")
        if item["completion"] == "UNKNOWN":
            reasons.append("COMPLETION_NOT_CONFIRMED")
        if item["completion"] == "COMPLETED" and item["credits"] is None:
            reasons.append("EARNED_CREDITS_MISSING")
        if item["completion"] == "COMPLETED" and candidate["raw_values"]["earned_credits"] is None:
            reasons.append("EARNED_CREDITS_NOT_IN_UPLOAD")
        if item["completion"] == "COMPLETED" and candidate["raw_values"]["semester"] is None:
            reasons.append("SEMESTER_NOT_IN_UPLOAD")
        if course and item["completion"] == "COMPLETED" and item["credits"] is not None and item["credits"] != course["catalog_credits"]:
            reasons.append("CREDIT_DIFFERS_FROM_CATALOG")
        corrections = {}
        for key, current in (("course_code", item["code"] or None), ("course_name", item["name"] or None),
                             ("earned_credits", item["credits"]), ("semester", item["semester"]),
                             ("classification", item["classification"])):
            original = candidate["raw_values"][key]
            if original is not None and current is not None and (
                    _compact(str(original)) != _compact(str(current)) if key in {"course_code", "course_name", "classification"}
                    else str(original) != str(current)):
                corrections[key] = {"original": original, "reviewed": current}
        if corrections:
            reasons.append("USER_CORRECTION_NEEDS_SOURCE_RECHECK")
        marker = candidate["raw_values"]["completion_marker"]
        if item["completion"] == "COMPLETED" and marker is None:
            reasons.append("COMPLETION_NOT_IN_UPLOAD")
        if marker in {"미이수", "미취득", "낙제", "수강중", "U", "F"} and item["completion"] == "COMPLETED":
            reasons.append("COMPLETION_CONTRADICTS_UPLOAD")
        if item["code"] and code_counts[item["code"]] > 1:
            reasons.append("DUPLICATE_OR_RETAKE_NEEDS_REVIEW")
        if course and item["classification"] and item["classification"] in classification_names.values():
            if item["classification"] != classification_names.get(course["classification"]):
                reasons.append("CLASSIFICATION_DIFFERS_FROM_CATALOG")
        if row.get("confirmed") is not True:
            reasons.append("USER_REVIEW_REQUIRED")
        verified = not reasons
        resolution = ("DUPLICATE_CANDIDATE" if "DUPLICATE_OR_RETAKE_NEEDS_REVIEW" in reasons else
                      "UNRESOLVED" if not course else "AMBIGUOUS" if
                      any(reason in reasons for reason in ("COURSE_NAME_CONFLICT", "CLASSIFICATION_DIFFERS_FROM_CATALOG"))
                      else "RESOLVED_TO_CATALOG")
        review_status = "VERIFIED_FROM_UPLOAD" if verified else "UNVERIFIED"
        record = {"candidate_id": candidate["candidate_id"], "raw_values": candidate["raw_values"],
                  "reviewed_values": {"course_code": item["code"] or None, "course_name": item["name"] or None,
                                      "earned_credits": item["credits"], "semester": item["semester"],
                                      "classification": item["classification"],
                                      "completion_status": item["completion"]},
                  "corrections": corrections,
                  "official_course": {"course_id": course["course_id"], "name": course["name"],
                                      "credits": course["catalog_credits"], "classification": course["classification"],
                                      "source": course["source"]} if course else None,
                  "resolution_status": resolution, "review_status": review_status,
                  "extraction_status": candidate["extraction_status"],
                  "confidence": candidate["confidence"], "ocr_confidence_score": candidate["ocr_confidence_score"],
                  "verification_status": "VERIFIED" if verified else "UNVERIFIED",
                  "usable_for_decision": verified, "needs_user_confirmation": reasons,
                  "source_location": candidate["location"]}
        records.append(record)
        if reasons:
            attention.append({"candidate_id": candidate["candidate_id"], "reasons": reasons})
        attempts.append({"attempt_id": candidate["candidate_id"],
                         "course_id": course["course_id"] if course else None,
                         "completion_status": item["completion"],
                         "verification_status": "VERIFIED" if verified else "UNVERIFIED",
                         "evidence_id": ("USER-REVIEWED-" if row.get("confirmed") is True else "USER-UPLOAD-")
                         + candidate["candidate_id"],
                         "earned_credits": item["credits"], "semester": item["semester"],
                         "source_grade": candidate["raw_values"].get("grade"),
                         "source_course_code": candidate["raw_values"]["course_code"],
                         "source_course_name": candidate["raw_values"]["course_name"],
                         "review_status": review_status, "resolution_status": resolution,
                         "evidence_type": ("USER_CONFIRMED_UPLOAD" if row.get("confirmed") is True
                                           else "USER_UPLOADED_CONTEXT"),
                         "evidence_locator": {"source_sha256": extracted["source_id"],
                                              "filename": extracted["filename"],
                                              "location": candidate["location"]}})
    fields = review.get("student_fields") or {}
    if not isinstance(fields, dict):
        raise ValueError("Reviewed student fields must be an object")
    reviewed_identity = fields.get("confirmed") is True
    student_id = str(fields.get("student_id") or "").strip() or None if reviewed_identity else None
    department = str(fields.get("department") or "").strip() if reviewed_identity else ""
    department_id = "DEPT-COMPUTER-ENGINEERING" if department in {"컴퓨터공학과", "DEPT-COMPUTER-ENGINEERING"} else None
    program_type = fields.get("program_type") if reviewed_identity and fields.get("program_type") in {"SINGLE", "MINOR", "DOUBLE"} else None

    def year(key: str) -> int | None:
        value = fields.get(key) if reviewed_identity else None
        if value in (None, ""):
            return None
        if not re.fullmatch(r"20\d{2}", str(value)):
            raise ValueError(f"Reviewed {key} must be a four-digit year")
        return int(value)

    state = {"student_state_id": "UPLOAD-" + extracted["source_id"][:12],
             "student_id": student_id, "admission_year": year("admission_year"),
             "department_id": department_id, "credit_policy_year": year("credit_policy_year"),
             "catalog_year": year("catalog_year"), "applicability_status": "UNVERIFIED",
             "program_type": program_type, "completion_coverage": "PARTIAL",
             "course_attempts": attempts, "upload_source_sha256": extracted["source_id"]}
    field_reviews = {}
    source_field_names = {"student_id": "student_id", "department_id": "department",
                          "admission_year": "admission_year"}
    for field, source_field in source_field_names.items():
        source = extracted["student_fields"][source_field]
        observations = source["observations"]
        reviewed = state[field]
        if field == "department_id":
            matches = [value["raw"] in {"컴퓨터공학과", "DEPT-COMPUTER-ENGINEERING"}
                       for value in observations]
            agrees = reviewed == "DEPT-COMPUTER-ENGINEERING" and bool(matches) and all(matches)
        else:
            agrees = reviewed is not None and bool(observations) and all(
                str(value["raw"]) == str(reviewed) for value in observations)
        field_reviews[field] = {"raw": observations, "reviewed": reviewed,
                                "status": "VERIFIED_FROM_UPLOAD" if reviewed_identity and agrees and source["status"] == "UNVERIFIED"
                                else "NEEDS_INFORMATION"}
        if field_reviews[field]["status"] != "VERIFIED_FROM_UPLOAD":
            attention.append({"field": field, "reasons": ["UPLOAD_IDENTITY_FIELD_NOT_CONFIRMED_BY_SOURCE"]})
    for field in ("credit_policy_year", "catalog_year", "program_type"):
        field_reviews[field] = {"raw": [], "reviewed": state[field], "status": "NEEDS_OFFICIAL_APPLICABILITY"}
    for key in ("student_id", "department_id", "admission_year", "credit_policy_year", "catalog_year", "program_type"):
        if state[key] is None:
            attention.append({"field": key, "reasons": ["STUDENT_FIELD_NEEDS_CONFIRMATION"]})
    if extracted["document_coverage"]["unreadable_pages"]:
        attention.append({"field": "document_coverage", "reasons": ["UNREADABLE_PAGES"]})
    attention.append({"field": "completion_coverage", "reasons": ["FULL_TRANSCRIPT_COVERAGE_NOT_VERIFIED"]})
    attention.append({"field": "applicability_status", "reasons": ["OFFICIAL_APPLICABILITY_NOT_VERIFIED"]})
    return {"source_id": extracted["source_id"], "student_state": state,
            "raw_student_fields": extracted["student_fields"], "student_field_reviews": field_reviews,
            "document_coverage": extracted["document_coverage"],
            "records": records, "needs_user_confirmation": attention,
            "usable_record_count": sum(r["usable_for_decision"] for r in records),
            "unresolved_record_count": sum(not r["usable_for_decision"] for r in records),
            "ruleset_id": catalog["curriculum_ruleset"]["ruleset_id"],
            "ruleset_version": catalog["curriculum_ruleset"]["ruleset_version"]}
