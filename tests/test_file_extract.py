"""Upload extraction tests, including reviewer counterexamples."""
from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))
from build_catalog import build  # noqa: E402
from curriculum_assistant.file_extract import extract_upload  # noqa: E402


def korean_image() -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1500, 260), "white")
    draw = ImageDraw.Draw(image)
    draw.text((45, 45), "CDA0143 고급자료구조 3학점", fill="black",
              font=ImageFont.truetype("C:/Windows/Fonts/malgun.ttf", 48))
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


class FileExtractionCases(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = build()

    def test_docx_table_yields_only_unverified_candidate(self):
        from docx import Document
        document = Document()
        table = document.add_table(rows=1, cols=3)
        for cell, value in zip(table.rows[0].cells, ("CDA0143", "고급자료구조", "3")):
            cell.text = value
        output = io.BytesIO()
        document.save(output)
        result = extract_upload("grade.docx", output.getvalue(), self.catalog)
        self.assertEqual(1, result["candidate_count"])
        candidate = result["candidates"][0]
        self.assertEqual("CDA0143", candidate["course_id"])
        self.assertEqual("VERIFIED", candidate["catalog_verification_status"])
        self.assertEqual("UNVERIFIED", candidate["extraction_status"])
        self.assertEqual("UNKNOWN", candidate["completion_status"])
        self.assertTrue(result["requires_user_confirmation"])

    def test_text_pdf_extracts_course_mention(self):
        from reportlab.pdfgen.canvas import Canvas
        output = io.BytesIO()
        canvas = Canvas(output)
        canvas.drawString(60, 750, "CDA0163 3")
        canvas.save()
        result = extract_upload("grade.pdf", output.getvalue(), self.catalog)
        self.assertEqual("CDA0163", result["candidates"][0]["course_id"])
        self.assertIn("PDF page 1", result["candidates"][0]["location"])

    def test_unknown_code_is_never_catalog_verified(self):
        from docx import Document
        document = Document()
        document.add_paragraph("XYZ9999 completed 3")
        output = io.BytesIO()
        document.save(output)
        candidate = extract_upload("grade.docx", output.getvalue(), self.catalog)["candidates"][0]
        self.assertEqual("MISSING", candidate["catalog_verification_status"])
        self.assertIsNone(candidate["catalog_credits"])

    def test_review_counterexample_word_completed_does_not_verify_attempt(self):
        from docx import Document
        document = Document()
        document.add_paragraph("CDA0143 completed with grade A")
        output = io.BytesIO()
        document.save(output)
        candidate = extract_upload("grade.docx", output.getvalue(), self.catalog)["candidates"][0]
        self.assertEqual("UNKNOWN", candidate["completion_status"])
        self.assertEqual("UNVERIFIED", candidate["extraction_status"])

    def test_bad_signature_rejected(self):
        with self.assertRaises(ValueError):
            extract_upload("grade.pdf", b"CDA0143", self.catalog)
        with self.assertRaises(ValueError):
            extract_upload("grade.hwp", b"CDA0143", self.catalog)

    def test_local_korean_ocr_returns_unverified_candidate(self):
        if not (ROOT / ".ocrmodels/korean_g2.pth").exists():
            self.skipTest("Optional local Korean OCR models are not installed")
        result = extract_upload("grade.png", korean_image(), self.catalog)
        candidate = result["candidates"][0]
        self.assertEqual("CDA0143", candidate["course_id"])
        self.assertEqual("UNVERIFIED", candidate["extraction_status"])
        self.assertEqual("UNKNOWN", candidate["completion_status"])
        self.assertGreater(candidate["ocr_confidence_score"], 0)

    def test_scanned_pdf_uses_local_ocr_candidate(self):
        if not (ROOT / ".ocrmodels/korean_g2.pth").exists():
            self.skipTest("Optional local Korean OCR models are not installed")
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen.canvas import Canvas
        output = io.BytesIO()
        canvas = Canvas(output, pagesize=(750, 240))
        canvas.drawImage(ImageReader(io.BytesIO(korean_image())), 0, 0, width=750, height=130)
        canvas.save()
        result = extract_upload("scanned.pdf", output.getvalue(), self.catalog)
        self.assertEqual("CDA0143", result["candidates"][0]["course_id"])
        self.assertIn("OCR", result["candidates"][0]["location"])


if __name__ == "__main__":
    unittest.main()
