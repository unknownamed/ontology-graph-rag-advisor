"""Adversarial transcript normalization tests using synthetic documents only."""
from __future__ import annotations

import io
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "src"), str(ROOT / "scripts")]
from build_catalog import build  # noqa: E402
from curriculum_assistant.file_extract import extract_upload, normalize_upload  # noqa: E402
from curriculum_assistant.engine import execute  # noqa: E402
from curriculum_assistant.graph import Graph  # noqa: E402


def transcript(rows: list[list[str]], *, header: str = "학번: 2026123456 | 학과: 컴퓨터공학과 | 입학연도: 2026") -> bytes:
    from docx import Document
    document = Document()
    document.add_paragraph(header)
    table = document.add_table(rows=0, cols=6)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def completed_review(result: dict, *, confirm: bool = True) -> dict:
    return {"student_fields": {"confirmed": True, "student_id": "2026123456",
                               "department": "컴퓨터공학과", "admission_year": "2026",
                               "credit_policy_year": "2026", "catalog_year": "2026",
                               "program_type": "SINGLE"},
            "records": [{"candidate_id": item["candidate_id"],
                         "course_code": item["raw_values"]["course_code"],
                         "course_name": item["raw_values"]["course_name"],
                         "earned_credits": item["raw_values"]["earned_credits"],
                         "semester": item["raw_values"]["semester"],
                         "classification": item["raw_values"]["classification"],
                         "completion_status": "COMPLETED", "confirmed": confirm}
                        for item in result["candidates"]]}


class StudentUploadNormalizationCases(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = build()

    def test_realistic_row_stays_unverified_until_explicit_review(self):
        data = transcript([["2026-1학기", "CDA0143", "고급자료구조", "3", "전공필수", "이수"]])
        extracted = extract_upload("TEST_FIXTURE.docx", data, self.catalog)
        self.assertEqual(1, extracted["candidate_count"])
        candidate = extracted["candidates"][0]
        self.assertEqual({"course_code": "CDA0143", "course_name": "고급자료구조", "earned_credits": 3,
                          "semester": "2026-1", "classification": "전공필수", "completion_marker": "이수"},
                         candidate["raw_values"])
        self.assertEqual("RESOLVED_TO_CATALOG", candidate["resolution_status"])
        draft = normalize_upload(extracted, {"records": []}, self.catalog)
        self.assertEqual("UNVERIFIED", draft["student_state"]["course_attempts"][0]["verification_status"])
        self.assertEqual("USER_UPLOADED_CONTEXT", draft["student_state"]["course_attempts"][0]["evidence_type"])
        self.assertTrue(draft["student_state"]["course_attempts"][0]["evidence_id"].startswith("USER-UPLOAD-"))
        self.assertIsNone(draft["student_state"]["student_id"])
        self.assertEqual("PARTIAL", draft["student_state"]["completion_coverage"])
        self.assertEqual("UNVERIFIED", draft["student_state"]["applicability_status"])
        reviewed = normalize_upload(extracted, completed_review(extracted), self.catalog)
        self.assertEqual("VERIFIED_FROM_UPLOAD", reviewed["records"][0]["review_status"])
        self.assertEqual("USER_CONFIRMED_UPLOAD", reviewed["student_state"]["course_attempts"][0]["evidence_type"])
        self.assertTrue(reviewed["records"][0]["usable_for_decision"])
        self.assertEqual("CDA0143", reviewed["student_state"]["course_attempts"][0]["course_id"])
        self.assertEqual(3, reviewed["student_state"]["course_attempts"][0]["earned_credits"])
        self.assertEqual("UNVERIFIED", reviewed["student_state"]["applicability_status"])
        self.assertEqual("CRS-CE-2026-CORE", reviewed["ruleset_id"])

    def test_name_disagreement_remains_ambiguous(self):
        data = transcript([["2026-1학기", "CDA0143", "다른과목", "3", "전공필수", "이수"]])
        extracted = extract_upload("TEST_FIXTURE.docx", data, self.catalog)
        self.assertEqual("AMBIGUOUS", extracted["candidates"][0]["resolution_status"])
        reviewed = normalize_upload(extracted, completed_review(extracted), self.catalog)
        self.assertEqual("AMBIGUOUS", reviewed["records"][0]["resolution_status"])
        self.assertFalse(reviewed["records"][0]["usable_for_decision"])
        self.assertIn("COURSE_NAME_CONFLICT", reviewed["records"][0]["needs_user_confirmation"])

    def test_correction_is_visible_but_not_silently_verified(self):
        data = transcript([["2026-1학기", "CDA0143", "고급자료구조", "2", "전공필수", "이수"]])
        extracted = extract_upload("TEST_FIXTURE.docx", data, self.catalog)
        review = completed_review(extracted)
        review["records"][0]["earned_credits"] = 3
        reviewed = normalize_upload(extracted, review, self.catalog)
        self.assertEqual(2, reviewed["records"][0]["raw_values"]["earned_credits"])
        self.assertEqual(3, reviewed["records"][0]["reviewed_values"]["earned_credits"])
        self.assertIn("USER_CORRECTION_NEEDS_SOURCE_RECHECK", reviewed["records"][0]["needs_user_confirmation"])
        self.assertFalse(reviewed["records"][0]["usable_for_decision"])

    def test_duplicate_course_rows_require_repeat_review(self):
        data = transcript([["2026-1학기", "CDA0143", "고급자료구조", "3", "전공필수", "이수"],
                           ["2026-2학기", "CDA0143", "고급자료구조", "3", "전공필수", "이수"]])
        extracted = extract_upload("TEST_FIXTURE.docx", data, self.catalog)
        self.assertEqual(["DUPLICATE_CANDIDATE"] * 2,
                         [c["resolution_status"] for c in extracted["candidates"]])
        reviewed = normalize_upload(extracted, completed_review(extracted), self.catalog)
        self.assertEqual(0, reviewed["usable_record_count"])
        self.assertTrue(all(r["verification_status"] == "UNVERIFIED" for r in reviewed["records"]))

    def test_unknown_code_and_name_only_never_gain_catalog_credit(self):
        data = transcript([["2026-1학기", "XYZ9999", "가상과목", "3", "전공선택", "이수"],
                           ["2026-2학기", "", "고급자료구조", "3", "전공필수", "이수"]])
        extracted = extract_upload("TEST_FIXTURE.docx", data, self.catalog)
        self.assertEqual(2, extracted["candidate_count"])
        self.assertEqual("UNRESOLVED", extracted["candidates"][0]["resolution_status"])
        self.assertEqual("AMBIGUOUS", extracted["candidates"][1]["resolution_status"])
        reviewed = normalize_upload(extracted, completed_review(extracted), self.catalog)
        self.assertEqual(0, reviewed["usable_record_count"])
        self.assertTrue(all(a["course_id"] is None for a in reviewed["student_state"]["course_attempts"]))

    def test_missing_credit_or_completion_marker_cannot_be_filled_from_catalog(self):
        data = transcript([["2026-1학기", "CDA0143", "고급자료구조", "", "전공필수", ""]])
        extracted = extract_upload("TEST_FIXTURE.docx", data, self.catalog)
        review = completed_review(extracted)
        review["records"][0]["earned_credits"] = 3
        reviewed = normalize_upload(extracted, review, self.catalog)
        self.assertEqual(0, reviewed["usable_record_count"])
        self.assertIn("EARNED_CREDITS_NOT_IN_UPLOAD", reviewed["records"][0]["needs_user_confirmation"])
        self.assertIn("COMPLETION_NOT_IN_UPLOAD", reviewed["records"][0]["needs_user_confirmation"])

    def test_truthy_text_is_not_user_confirmation(self):
        data = transcript([["2026-1학기", "CDA0143", "고급자료구조", "3", "전공필수", "이수"]])
        extracted = extract_upload("TEST_FIXTURE.docx", data, self.catalog)
        review = completed_review(extracted)
        review["records"][0]["confirmed"] = "true"
        review["student_fields"]["confirmed"] = "true"
        normalized = normalize_upload(extracted, review, self.catalog)
        self.assertEqual(0, normalized["usable_record_count"])
        self.assertIsNone(normalized["student_state"]["student_id"])

    def test_unresolved_completion_remains_information_gap_when_applied_later(self):
        student = json.loads((ROOT / "tests/fixtures/core_eligible_synthetic.json").read_text(
            encoding="utf-8"))["student_state"]
        student["course_attempts"].append({"attempt_id": "TEST-FIXTURE-UNKNOWN-COMPLETION",
                                            "course_id": "CDA0143", "completion_status": "UNKNOWN",
                                            "verification_status": "UNVERIFIED",
                                            "evidence_id": "TEST-FIXTURE-UPLOAD", "earned_credits": None})
        graph = Graph(Path(":memory:"), self.catalog)
        try:
            decision = execute(graph, student, {"intent": "GRADUATION_STATUS"})["decision"]
        finally:
            graph.close()
        self.assertEqual("UNKNOWN", decision["graduation_outcome"])
        self.assertEqual("NEEDS_INFORMATION", decision["decision_status"])
        self.assertTrue(any("COMPLETION_STATUS:TEST-FIXTURE-UNKNOWN-COMPLETION" in value
                            for value in decision["needs_information"]))

    def test_blank_pdf_page_is_reported_and_full_coverage_remains_unknown(self):
        from reportlab.pdfgen.canvas import Canvas
        output = io.BytesIO()
        canvas = Canvas(output)
        canvas.drawString(60, 750, "CDA0143 3")
        canvas.showPage()
        canvas.drawString(60, 750, " ")
        canvas.save()
        with patch("curriculum_assistant.file_extract._ocr_image", return_value=([], [], {})):
            extracted = extract_upload("TEST_FIXTURE.pdf", output.getvalue(), self.catalog)
        self.assertEqual([2], extracted["document_coverage"]["unreadable_pages"])
        self.assertIsNone(extracted["document_coverage"]["missing_pages"])
        self.assertEqual("UNKNOWN", extracted["document_coverage"]["completion_coverage"])

    def test_no_real_student_defaults_are_invented(self):
        data = transcript([["2026-1학기", "CDA0143", "고급자료구조", "3", "전공필수", "이수"]], header="합성 성적표")
        extracted = extract_upload("TEST_FIXTURE.docx", data, self.catalog)
        reviewed = normalize_upload(extracted, {"records": []}, self.catalog)
        state = reviewed["student_state"]
        self.assertIsNone(state["department_id"])
        self.assertIsNone(state["admission_year"])
        self.assertIsNone(state["catalog_year"])
        self.assertIsNone(state["program_type"])
        self.assertFalse(any(a["verification_status"] == "VERIFIED" for a in state["course_attempts"]))

    def test_pdf_table_columns_are_not_lost_or_double_counted(self):
        from reportlab.lib import colors
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.platypus import Table, TableStyle

        output = io.BytesIO()
        canvas = Canvas(output)
        for row in (["2026", "1", "", "CDA0143", "TestCourse", "3", "", "A+"],
                    ["2026", "2", "", "CDA0157", "AnotherCourse", "3", "", "B0"]):
            table = Table([["Year", "Term", "Type", "Code", "Name", "Credits", "Design", "Grade"], row],
                          colWidths=[52, 46, 44, 72, 110, 52, 46, 48])
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
            table.wrapOn(canvas, 0, 0)
            table.drawOn(canvas, 40, 650)
            canvas.showPage()
        canvas.save()
        extracted = extract_upload("TEST_FIXTURE.pdf", output.getvalue(), self.catalog)
        self.assertEqual(2, extracted["document_coverage"]["page_count"])
        self.assertEqual(2, extracted["candidate_count"])
        self.assertEqual(["PDF page 1, table 1, row 2", "PDF page 2, table 1, row 2"],
                         [row["location"] for row in extracted["candidates"]])
        self.assertEqual([("TestCourse", 3, "2026-1", "A+"),
                          ("AnotherCourse", 3, "2026-2", "B0")],
                         [(row["raw_values"]["course_name"], row["raw_values"]["earned_credits"],
                           row["raw_values"]["semester"], row["raw_values"]["grade"])
                          for row in extracted["candidates"]])

    def test_grade_u_and_winter_term_remain_explicit_candidates(self):
        data = transcript([["2025-WINTER", "CDA0143", "고급자료구조", "0", "전공필수", "U"]])
        extracted = extract_upload("TEST_FIXTURE.docx", data, self.catalog)
        raw = extracted["candidates"][0]["raw_values"]
        self.assertEqual("2025-WINTER", raw["semester"])
        self.assertEqual("U", raw["grade"])
        self.assertEqual("U", raw["completion_marker"])
        review = completed_review(extracted)
        normalized = normalize_upload(extracted, review, self.catalog)
        self.assertIn("COMPLETION_CONTRADICTS_UPLOAD", normalized["records"][0]["needs_user_confirmation"])
        self.assertFalse(normalized["records"][0]["usable_for_decision"])


if __name__ == "__main__":
    unittest.main()
